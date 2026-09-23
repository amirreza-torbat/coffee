# -*- coding: utf-8 -*-
"""Build the Caspiva outreach Word document: 163 ACTIVE GREEN COFFEE EXPORTERS,
one company per page, SL No shown on every page, copy-ready WhatsApp message."""
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = 'Indian Coffee Exporters - Complete Database (951).xlsx'
OUT = 'Caspiva - Green Coffee Outreach (163 companies).docx'

wb = load_workbook(SRC, read_only=True)
ws = wb['A-GREEN COFFEE']
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
I = {h: i for i, h in enumerate(hdr)}
rows = sorted(ws.iter_rows(min_row=2, values_only=True), key=lambda r: r[I['SL No']])
def g(r, k):
    v = r[I[k]]
    return '' if v is None else str(v).strip()

# ---------------------------------------------------------------- helpers
def shade(el_pr, color):
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), color)
    el_pr.append(shd)
def para_shade(p, color): shade(p._p.get_or_add_pPr(), color)
def rtl(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:bidi'); b.set(qn('w:val'), '1'); pPr.append(b)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
def field(par, code):
    r = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r._r.append(fc)
    r2 = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = code; r2._r.append(it)
    r3 = par.add_run(); fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'separate'); r3._r.append(fc2)
    r4 = par.add_run('1')
    r5 = par.add_run(); fc3 = OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'), 'end'); r5._r.append(fc3)
    return par

# ---------------------------------------------------------------- message template
# (text, bold?)  -- bold segments are also wrapped in WhatsApp '*' markers
def msg_paras(name):
    return [
        [("Dear Sales & Export Team at ", 0), (name, 1)],
        [("Good day.", 0)],
        [("This is the Procurement Team from ", 0), ("Caspiva", 1),
         (", an importer and distributor of green coffee beans in Iran.", 0)],
        [("We are sourcing upcoming FCL orders and would like to request your ", 0),
         ("latest product catalog, technical spec sheets, and price list", 1), (" for:", 0)],
        [("1. ", 0), ("India Robusta:", 1),
         (" Cherry AA (Scr 18), Cherry AB (Scr 15/16), Kaapi Royale (Scr 18), and ", 0),
         ("Robusta PB (Peaberry)", 1), (".", 0)],
        [("2. ", 0), ("India Arabica:", 1),
         (" Plantation AA (Scr 18), Plantation A (Scr 17), and ", 0), ("Plantation PB", 1), (".", 0)],
        [("3. ", 0), ("Vietnam Origins (if available):", 1),
         (" Robusta Grade 1 (Scr 18 & 16 Wet Polished).", 0)],
        [("Requirements:", 1)],
        [("- Current Crop, Color Sorted, Moisture , Black beans , Broken .", 0)],
        [("- 60kg Jute bags with GrainPro liner.", 0)],
        [("Please kindly provide:", 1)],
        [("- ", 0), ("FOB (Indian ports) & CIF Chabahar / CIF Khasab", 1), (" prices.", 0)],
        [("- ", 0), ("Payment terms & Lead time", 1), (" from order confirmation.", 0)],
        [("- Sample dispatch procedure (500g pre-shipment green samples).", 0)],
        [("Looking forward to your prompt feedback.", 0)],
        [("Best regards,", 0)],
        [("Caspiva co.", 1)],
        [("WhatsApp: +989058406009", 0)],
    ]

# ---------------------------------------------------------------- document
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
sec = doc.sections[0]
sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.4); sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)

# running header
hp = sec.header.paragraphs[0]
hp.text = ''
r = hp.add_run('ACTIVE GREEN COFFEE EXPORTERS  •  163 companies  •  Caspiva procurement outreach  •  Coffee Board of India database (2023)')
r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x59, 0x59, 0x59); r.font.bold = True
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
# footer with page numbers
fp = sec.footer.paragraphs[0]; fp.text = ''
fr = fp.add_run('Page '); fr.font.size = Pt(8)
field(fp, ' PAGE ')
fr2 = fp.add_run(' of '); fr2.font.size = Pt(8)
field(fp, ' NUMPAGES ')
fr3 = fp.add_run('  —  هر صفحه = یک شرکت  —  شماره ردیف در بالای هر صفحه آمده است')
fr3.font.size = Pt(8); rtl(fp)

# ---- title + instructions
t = doc.add_paragraph(); tr = t.add_run('پیام استعلام قیمت — صادرکنندگان فعال قهوه سبز هند (Category A)')
tr.font.size = Pt(16); tr.font.bold = True; tr.font.color.rgb = RGBColor(0x00, 0x61, 0x00); rtl(t)
p = doc.add_paragraph(); rtl(p)
p.add_run('این فایل ۱۶۳ صفحه پیام آماده دارد: هر صفحه مربوط به یک شرکت است و شماره ردیف (SL No) آن شرکت در نوار خاکستری بالای همان صفحه نوشته شده تا هیچ اطلاعاتی گم نشود. '
          'فقط متن داخل کادر را کپی کنید و برای شماره واتس‌اپ همان شرکت بفرستید؛ نوار خاکستری بالا فقط برای مرجع شماست و جزو پیام نیست. '
          'علامت‌های * داخل متن، قالب‌بندی خودِ واتس‌اپ (بولد) هستند و پس از ارسال به‌درستی نمایش داده می‌شوند.').font.size = Pt(10)
p2 = doc.add_paragraph(); rtl(p2)
p2.add_run('شرکت‌هایی که در منبع رسمی شماره موبایل ندارند (۲۴ مورد) در نوار خاکستری با عبارت «NO MOBILE» علامت خورده‌اند؛ پیام آن‌ها را با ایمیل یا تلفن ثابت پیگیری کنید.').font.size = Pt(10)

# ---- index table
doc.add_paragraph()
h = doc.add_paragraph(); hr_ = h.add_run('فهرست سریع (جهت ارجاع)'); hr_.font.bold = True; hr_.font.size = Pt(12); rtl(h)
tbl = doc.add_table(rows=1, cols=5); tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hc = ['SL No', 'Company Name', 'Mobile / WhatsApp', 'Landline', 'E-mail']
for j, txt in enumerate(hc):
    c = tbl.rows[0].cells[j]; c.text = ''
    rr = c.paragraphs[0].add_run(txt); rr.font.bold = True; rr.font.size = Pt(9)
    shade(c._tc.get_or_add_tcPr(), 'C6EFCE')
for r_ in rows:
    cells = tbl.add_row().cells
    vals = [g(r_, 'SL No'), g(r_, 'Company Name'),
            g(r_, 'All Mobiles') if g(r_, 'All Mobiles') not in ('', '—') else 'NO MOBILE',
            g(r_, 'Landline(s)') if g(r_, 'Landline(s)') not in ('', '—') else '—',
            g(r_, 'All E-mails') if g(r_, 'All E-mails') not in ('', '—') else '—']
    for j, v in enumerate(vals):
        cells[j].text = ''
        rr = cells[j].paragraphs[0].add_run(v); rr.font.size = Pt(8)
        if j == 0: rr.font.bold = True
        if j == 2 and v == 'NO MOBILE': rr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00); rr.font.bold = True
for j, w in enumerate([Cm(1.3), Cm(6.2), Cm(3.6), Cm(4.2), Cm(4.6)]):
    for row in tbl.rows: row.cells[j].width = w

# ---- one page per company
first = True
for r_ in rows:
    sl = g(r_, 'SL No'); name = g(r_, 'Company Name')
    mobs = [m for m in g(r_, 'All Mobiles').split('; ') if m and m != '—']
    lls  = [m for m in g(r_, 'Landline(s)').split('; ') if m and m != '—']
    ems  = [m for m in g(r_, 'All E-mails').split('; ') if m and m != '—']
    if not first: doc.add_page_break()
    first = False

    hp_ = doc.add_paragraph(); para_shade(hp_, 'D9D9D9')
    hp_.paragraph_format.space_before = Pt(2); hp_.paragraph_format.space_after = Pt(4)
    a = hp_.add_run('SL No %s / 163' % sl); a.font.bold = True; a.font.size = Pt(11); a.font.color.rgb = RGBColor(0x00, 0x61, 0x00)
    b = hp_.add_run('   |   %s' % name); b.font.bold = True; b.font.size = Pt(11)
    if mobs:
        c1 = hp_.add_run('   |   WhatsApp: %s' % ' , '.join(mobs)); c1.font.size = Pt(10); c1.font.bold = True
        c1.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    else:
        c1 = hp_.add_run('   |   NO MOBILE in source'); c1.font.size = Pt(10); c1.font.bold = True
        c1.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    if lls:
        c2 = hp_.add_run('   |   Tel: %s' % ' , '.join(lls)); c2.font.size = Pt(9)
    if ems:
        c3 = hp_.add_run('   |   E-mail: %s' % ' , '.join(ems)); c3.font.size = Pt(9)
        c3.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    hint = doc.add_paragraph()
    hh = hint.add_run('▼ متن زیر را کپی کنید (COPY FROM HERE) ▼'); hh.font.size = Pt(8); hh.font.italic = True
    hh.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    hint.paragraph_format.space_after = Pt(2)

    box = doc.add_table(rows=1, cols=1); box.style = 'Table Grid'
    cell = box.rows[0].cells[0]
    shade(cell._tc.get_or_add_tcPr(), 'FFFFFF')
    cell.text = ''
    paras = msg_paras(name)
    for pi, segs in enumerate(paras):
        par = cell.paragraphs[0] if pi == 0 else cell.add_paragraph()
        par.paragraph_format.space_after = Pt(4); par.paragraph_format.space_before = Pt(0)
        for txt, bold in segs:
            run = par.add_run(('*%s*' % txt) if bold else txt)
            run.font.size = Pt(10.5)
            if bold: run.font.bold = True
    end = doc.add_paragraph()
    ee = end.add_run('▲ پایان متن این شرکت — شرکت بعدی در صفحه بعد ▲'); ee.font.size = Pt(8); ee.font.italic = True
    ee.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

doc.save(OUT)
print('saved', OUT, '| companies:', len(rows))
